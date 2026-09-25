# -*- coding: utf-8 -*-
"""Baut die Seminarpräsentation 'Erfolgsfaktor Emotionale Intelligenz' auf Basis der Best-Akademie-Vorlage."""
import copy, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
SEMINAR = "Erfolgsfaktor Emotionale Intelligenz"

BLUE, TEAL, DARK = "0B5394", "45818E", "1F2A36"
LBLUE, LTEAL, GREY = "DCE8F3", "DDEDEF", "F2F4F6"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
A = "{%s}" % NS["a"]

prs = Presentation(TEMPLATE)
tpl_slides = list(prs.slides)
title_slide, proto, end_slide = tpl_slides[0], tpl_slides[1], tpl_slides[-1]
content_layout = proto.slide_layout

proto_tree = proto.shapes._spTree
proto_sps = {sp.name: sp._element for sp in proto.shapes}
P_TITLE = proto.shapes[0]._element
P_BODY = proto.shapes[1]._element
P_NUM = proto.shapes[2]._element
P_SRC = proto.shapes[3]._element
P_BG = proto._element.find("p:cSld/p:bg", NS)

# ---------------------------------------------------------------- Textbausteine
FONT = '<a:latin typeface="Arial"/><a:ea typeface="Arial"/><a:cs typeface="Arial"/><a:sym typeface="Arial"/>'


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def rpr(sz, color=None, b=False, i=False, scheme=False):
    fill = '<a:solidFill><a:schemeClr val="dk1"/></a:solidFill>' if (scheme or color is None) else \
        '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % color
    return '<a:rPr lang="de-DE" sz="%d"%s%s>%s%s</a:rPr>' % (
        sz, ' b="1"' if b else "", ' i="1"' if i else "", fill, FONT)


def runs(text, sz, color=None, b=False, i=False):
    """**fett** im Fließtext"""
    out = []
    for k, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        bold = b or (k % 2 == 1)
        out.append('<a:r>%s<a:t xml:space="preserve">%s</a:t></a:r>' % (rpr(sz, color, bold, i), esc(part)))
    return "".join(out)


def ppr(before, marL=0, indent=0, lvl=0, bu=None, after=0):
    if bu is None:
        b = '<a:buNone/>'
    elif bu == "num":
        b = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buSzPts val="1600"/><a:buFont typeface="Arial"/><a:buAutoNum type="arabicPeriod"/>' % TEAL
    elif bu == "●":
        b = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buSzPts val="1600"/><a:buFont typeface="Arial"/><a:buChar char="●"/>' % TEAL
    else:
        b = '<a:buClr><a:schemeClr val="dk1"/></a:buClr><a:buSzPts val="1600"/><a:buFont typeface="Arial"/><a:buChar char="○"/>'
    return ('<a:pPr indent="%d" lvl="%d" marL="%d" rtl="0" algn="l"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc>'
            '<a:spcBef><a:spcPts val="%d"/></a:spcBef><a:spcAft><a:spcPts val="%d"/></a:spcAft>%s</a:pPr>'
            % (indent, lvl, marL, before, after, b))


def para(kind, text, sz=None):
    if kind == "h":      # Kapitelzeile (wie Vorlage)
        return '<a:p>%s%s</a:p>' % (ppr(0), runs(text, sz or 1900, BLUE, b=True))
    if kind == "sh":     # Zwischenüberschrift
        return '<a:p>%s%s</a:p>' % (ppr(1400), runs(text, sz or 1700, BLUE, b=True))
    if kind == "p":
        return '<a:p>%s%s</a:p>' % (ppr(1000), runs(text, sz or 1600))
    if kind == "i":
        return '<a:p>%s%s</a:p>' % (ppr(1000), runs(text, sz or 1600, i=True))
    if kind in ("b", "n", "b2"):
        s = sz or 1600
        if kind == "b2":
            pp = ppr(300, 914400, -336550, 1, "○")
        else:
            pp = ppr(700, 457200, -336550, 0, "num" if kind == "n" else "●")
        if "::" in text:
            lab, rest = text.split("::", 1)
            body = runs(lab + ":", s, BLUE, b=True) + runs(rest, s)
        else:
            body = runs(text, s)
        return '<a:p>%s%s</a:p>' % (pp, body)
    raise ValueError(kind)


def set_body(sp_el, items):
    txBody = sp_el.find("p:txBody", NS)
    for p in txBody.findall("a:p", NS):
        txBody.remove(p)
    xml = "".join(para(*it) if isinstance(it, tuple) else para("p", it) for it in items)
    frag = etree.fromstring('<x xmlns:a="%s">%s</x>' % (NS["a"], xml))
    for p in frag:
        txBody.append(p)


def set_single_text(sp_el, text):
    """ersetzt den Text des ersten Runs, entfernt weitere Runs"""
    ps = sp_el.findall(".//a:p", NS)
    keep = next((p for p in ps if any((t.text or "").strip() for t in p.iter(A + "t"))), ps[0])
    rs = keep.findall("a:r", NS)
    rs[0].find("a:t", NS).text = text
    for r in rs[1:]:
        keep.remove(r)
    for p in ps:
        if p is not keep:
            p.getparent().remove(p)
    return rs[0]


# ---------------------------------------------------------------- Folienfabrik
def new_slide(items, source="eigene Darstellung", body_h=None, notes=None):
    s = prs.slides.add_slide(content_layout)
    tree = s.shapes._spTree
    for sp in list(s.shapes):
        tree.remove(sp._element)
    bg = copy.deepcopy(P_BG)
    s._element.find("p:cSld", NS).insert(0, bg)
    t, b, n, q = (copy.deepcopy(e) for e in (P_TITLE, P_BODY, P_NUM, P_SRC))
    for e in (t, b, n, q):
        tree.append(e)
    set_single_text(t, SEMINAR).find("a:rPr", NS).set("sz", "3000")
    set_body(b, items)
    if body_h is not None:
        b.find("p:spPr/a:xfrm/a:ext", NS).set("cy", str(int(Inches(body_h))))
    set_single_text(q, "Quelle: " + source)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def hexrgb(h):
    return RGBColor.from_string(h)


def _nostyle(sh):
    st = sh._element.find("p:style", NS)
    if st is not None:
        sh._element.remove(st)


def chev(s, x, y, w, h, fill, lines, first=False):
    """Chevron/Pfeil ohne Text + Textbox darüber (verhindert Silbentrennung durch Formeinzug)"""
    box(s, x, y, w, h, fill, [], shape=MSO_SHAPE.PENTAGON if first else MSO_SHAPE.CHEVRON)
    ind = 0.12 if first else h * 0.45
    return box(s, x + ind, y, w - ind - h * 0.4, h, None, lines, margin=0.02)


def box(s, x, y, w, h, fill=BLUE, lines=(), size=14, color="FFFFFF", bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=None, margin=0.08):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = hexrgb(fill)
    if line:
        sh.line.color.rgb = hexrgb(line)
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    _nostyle(sh)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(margin + 0.04))
    tf.margin_top = tf.margin_bottom = Inches(margin)
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for ln in lines:
        if isinstance(ln, str):
            ln = (ln, size, bold, color)
        txt, sz, bd, col = ln
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        for k, part in enumerate(re.split(r"\*\*", txt)):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            f = r.font
            f.size, f.bold, f.name = Pt(sz), bd or (k % 2 == 1), "Arial"
            f.color.rgb = hexrgb(col)
    return sh


def arrow(s, x, y, w, h, fill=TEAL, shape=MSO_SHAPE.RIGHT_ARROW):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = hexrgb(fill)
    sh.line.fill.background()
    _nostyle(sh)
    return sh


def table(s, x, y, w, rows, colw, size=12, rowh=0.4, head_fill=BLUE, first_col_bold=True):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    tb = gt.table
    for j, cw in enumerate(colw):
        tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = hexrgb(head_fill if i == 0 else (LBLUE if i % 2 == 0 else "FFFFFF"))
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            for k, part in enumerate(re.split(r"\*\*", val)):
                if not part:
                    continue
                r = p.add_run()
                r.text = part
                r.font.size = Pt(size)
                r.font.name = "Arial"
                r.font.bold = (i == 0) or (j == 0 and first_col_bold) or (k % 2 == 1)
                r.font.color.rgb = hexrgb("FFFFFF" if i == 0 else (BLUE if j == 0 and first_col_bold else DARK))
    # Tabellenstil ohne Banding-Effekte
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    return gt


def divider(num, name, questions, source="eigene Darstellung", notes=None):
    s = new_slide([("h", "Modul %d" % num)], source, body_h=0.6, notes=notes)
    box(s, 0.71, 2.2, 5.6, 3.7, BLUE, [("MODUL %d" % num, 16, True, "C9DAF0"), (name, 30, True, "FFFFFF")],
        align=PP_ALIGN.LEFT, margin=0.35)
    box(s, 6.75, 2.2, 5.85, 0.55, None, [("Leitfragen", 17, True, BLUE)], align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.02)
    for k, q in enumerate(questions):
        y = 2.9 + k * 1.0
        box(s, 6.75, y, 0.6, 0.6, TEAL, [(str(k + 1), 18, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
        box(s, 7.55, y - 0.12, 5.05, 0.85, None, [(q, 15, False, DARK)], align=PP_ALIGN.LEFT, margin=0.02)
    return s


def exercise(title, ziel, schritte, dauer, form, source="eigene Darstellung", notes=None):
    s = new_slide([("h", "Übung: " + title), ("i", ziel)], source, body_h=1.3, notes=notes)
    y0 = 2.75
    for k, st in enumerate(schritte):
        y = y0 + k * 0.78
        box(s, 0.8, y, 0.55, 0.55, TEAL, [(str(k + 1), 16, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
        box(s, 1.5, y - 0.1, 6.9, 0.75, None, [(st, 15, False, DARK)], align=PP_ALIGN.LEFT, margin=0.02)
    box(s, 8.9, 2.75, 3.7, 3.3, LBLUE, [("Rahmen", 16, True, BLUE), ("", 6, False, DARK),
                                        ("Dauer: " + dauer, 14, False, DARK), ("", 6, False, DARK),
                                        ("Form: " + form, 14, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.25)
    return s


# ================================================================= INHALT
# 1 Titelfolie
set_single_text(title_slide.shapes[0]._element, SEMINAR)
title_slide.notes_slide.notes_text_frame.text = (
    "Begrüßung, technische Hinweise (Kamera an, Mikro stumm, Chat nutzen), kurzer Überblick über den Tag.")

# 2 Agenda
s = new_slide([("h", "Agenda – Online-Seminar (1 Tag)")], body_h=0.6)
table(s, 0.8, 2.05, 11.7, [
    ["Uhrzeit", "Inhalt", "Methodik"],
    ["09:00", "Begrüßung, Ziele, Erwartungen", "Plenum, Chat-Umfrage"],
    ["09:30", "Modul 1–2: Was ist Emotionale Intelligenz? Das Goleman-Modell", "Impuls, Diskussion"],
    ["10:30", "Pause", ""],
    ["10:45", "Modul 3: Selbstwahrnehmung", "Impuls, Selbstreflexion"],
    ["11:30", "Modul 4: Selbstregulation", "Impuls, Übung in Breakout-Räumen"],
    ["12:30", "Mittagspause", ""],
    ["13:15", "Modul 5: Motivation", "Impuls, Einzelarbeit"],
    ["14:00", "Modul 6: Empathie", "Impuls, Rollenspiel"],
    ["14:45", "Pause", ""],
    ["15:00", "Modul 7: Soziale Kompetenz", "Impuls, Fallarbeit"],
    ["15:45", "Modul 8: EI in Führung und Vertrieb", "Impuls, Praxisfälle"],
    ["16:30", "Modul 9: Transfer, Feedback, Abschluss", "Transferplan, Blitzlicht"],
], [1.3, 6.6, 3.8], size=12, rowh=0.36)

# 3 Seminarziele
s = new_slide([("h", "Seminarziele"), ("p", "Nach diesem Seminar können Sie …")], body_h=1.2)
goals = [("Verstehen", "Modelle und Kernkompetenzen der Emotionalen Intelligenz erklären"),
         ("Erkennen", "eigene Emotionen, Auslöser und Muster bewusst wahrnehmen"),
         ("Steuern", "Emotionen in Drucksituationen gezielt regulieren"),
         ("Verbinden", "Emotionen anderer erkennen und empathisch reagieren"),
         ("Wirken", "Gespräche, Feedback und Konflikte emotional intelligent führen"),
         ("Umsetzen", "einen persönlichen Transferplan für den Alltag erstellen")]
for k, (h, t) in enumerate(goals):
    x = 0.8 + (k % 3) * 3.95
    y = 2.6 + (k // 3) * 1.95
    box(s, x, y, 3.7, 1.7, LBLUE if k % 2 == 0 else LTEAL, [(h, 18, True, BLUE), (t, 14, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

# 4 Vorstellungsrunde
exercise("Vorstellungsrunde", "Ankommen, Kennenlernen und die eigenen Erwartungen sichtbar machen.",
         ["Name, Funktion und Unternehmen",
          "Eine Situation, in der Emotionen im Job eine große Rolle gespielt haben",
          "Meine Erwartung an den heutigen Tag (bitte zusätzlich in den Chat)",
          "Auf einer Skala von 1–10: Wie gut kenne ich meine emotionalen Auslöser?"],
         "ca. 2 Min. pro Person", "Plenum, Kamera an, Chat",
         notes="Erwartungen im Chat sammeln und am Ende des Tages wieder aufgreifen.")

# ---- Modul 1
divider(1, "Was ist Emotionale Intelligenz?",
        ["Wie wird Emotionale Intelligenz definiert?", "Worin unterscheiden sich IQ und EQ?",
         "Welche Modelle gibt es – und wie belastbar ist die Forschung?"])

s = new_slide([("h", "Definition: Emotionale Intelligenz"),
               ("i", "„Die Fähigkeit, die eigenen Gefühle und die anderer zu beobachten, zwischen ihnen zu unterscheiden "
                     "und diese Informationen zu nutzen, um das eigene Denken und Handeln zu steuern.“"),
               ("p", "– Peter Salovey & John D. Mayer (1990)"),
               ("sh", "Kurz gesagt"),
               ("b", "Emotionen **wahrnehmen**:: bei sich selbst und bei anderen"),
               ("b", "Emotionen **verstehen**:: Ursachen, Verläufe und Wirkungen erkennen"),
               ("b", "Emotionen **nutzen und steuern**:: für gute Entscheidungen und Beziehungen"),
               ("p", "Popularisiert wurde der Begriff 1995 durch **Daniel Goleman** mit dem Bestseller „Emotional Intelligence“ "
                     "(dt. „EQ. Emotionale Intelligenz“).")],
              "Salovey & Mayer (1990); Goleman (1995)")

s = new_slide([("h", "Die Geschichte des Konzepts")], "Thorndike (1920); Gardner (1983); Salovey & Mayer (1990); Goleman (1995); Bar-On (1997)",
              body_h=0.6)
tl = [("1920", "Edward Thorndike", "„Soziale Intelligenz“: die Fähigkeit, Menschen zu verstehen und klug mit ihnen umzugehen"),
      ("1983", "Howard Gardner", "Multiple Intelligenzen – u. a. inter- und intrapersonale Intelligenz"),
      ("1990", "Salovey & Mayer", "Erste wissenschaftliche Definition von Emotionaler Intelligenz"),
      ("1995", "Daniel Goleman", "Bestseller macht EI weltweit bekannt – Fokus auf Beruf und Führung"),
      ("1997", "Reuven Bar-On", "EQ-i: erster standardisierter Selbstauskunfts-Fragebogen")]
arrow(s, 0.8, 3.02, 11.8, 0.16, "B7C9D6", MSO_SHAPE.RECTANGLE)
for k, (yr, who, what) in enumerate(tl):
    x = 0.8 + k * 2.4
    box(s, x + 0.55, 2.6, 1.0, 1.0, BLUE if k % 2 == 0 else TEAL, [(yr, 15, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, x, 3.85, 2.2, 2.4, GREY, [(who, 14, True, BLUE), (what, 12, False, DARK)], align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.12)

s = new_slide([("h", "IQ und EQ im Vergleich")], "Goleman (1995); Mayer, Salovey & Caruso (2004)", body_h=0.6)
table(s, 0.8, 2.1, 11.7, [
    ["Merkmal", "Intelligenzquotient (IQ)", "Emotionale Intelligenz (EQ / EI)"],
    ["Beschreibt", "Kognitive Fähigkeiten: Logik, Sprache, Abstraktion", "Umgang mit eigenen und fremden Emotionen"],
    ["Veränderbarkeit", "Im Erwachsenenalter relativ stabil", "Durch Übung und Reflexion gut entwickelbar"],
    ["Messung", "Standardisierte Leistungstests", "Leistungstests (z. B. MSCEIT) oder Selbst-/Fremdeinschätzung"],
    ["Wirkt vor allem auf", "Fachliche Problemlösung, Lernen", "Zusammenarbeit, Führung, Stressbewältigung"],
    ["Typische Frage", "„Kann ich das Problem lösen?“", "„Wie gehe ich mit mir und den Beteiligten um?“"],
], [2.6, 4.4, 4.7], size=14, rowh=0.62)
box(s, 0.8, 6.0, 11.7, 0.5, LTEAL, [("IQ und EQ sind kein Gegensatz: Fachkompetenz öffnet Türen – Emotionale Intelligenz entscheidet oft, wie weit man kommt.", 14, True, BLUE)], margin=0.05)

s = new_slide([("h", "Drei Modelle der Emotionalen Intelligenz")], "Mayer & Salovey (1997); Goleman (1995, 1998); Petrides & Furnham (2001)", body_h=0.6)
mods = [("Fähigkeitsmodell", "Mayer & Salovey", "EI als **kognitive Fähigkeit**, Emotionen zu verarbeiten.",
         "Messung per Leistungstest (MSCEIT) – es gibt richtige und falsche Antworten."),
        ("Mischmodell", "Goleman / Bar-On", "EI als Bündel aus **Kompetenzen, Eigenschaften und Motivation**.",
         "Messung per Selbst- und 360°-Einschätzung (z. B. ESCI, EQ-i). Stark praxis- und führungsorientiert."),
        ("Trait-Modell", "Petrides & Furnham", "EI als **Persönlichkeitsmerkmal**: emotionale Selbstwirksamkeit.",
         "Messung per Selbstauskunft (TEIQue). Eng mit Big-Five-Persönlichkeit verbunden.")]
for k, (n, a, d, m) in enumerate(mods):
    x = 0.8 + k * 3.95
    box(s, x, 2.15, 3.7, 1.0, BLUE if k != 1 else TEAL, [(n, 19, True, "FFFFFF"), (a, 13, False, "E6EEF7")], margin=0.08)
    box(s, x, 3.25, 3.7, 3.0, GREY, [(d, 14, False, DARK), ("", 8, False, DARK), (m, 13, False, "4A5561")],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Das Vier-Zweige-Modell nach Mayer & Salovey"),
               ("p", "Die vier Fähigkeiten bauen aufeinander auf – von der einfachen Wahrnehmung bis zur bewussten Steuerung.")],
              "Mayer & Salovey (1997): What is emotional intelligence?", body_h=1.2)
br = [("1", "Wahrnehmen", "Emotionen in Mimik, Stimme, Körper und bei sich selbst erkennen"),
      ("2", "Nutzen", "Emotionen gezielt einsetzen, um Denken und Kreativität zu fördern"),
      ("3", "Verstehen", "Ursachen, Mischungen und Verläufe von Emotionen begreifen"),
      ("4", "Regulieren", "Emotionen bei sich und anderen konstruktiv beeinflussen")]
for k, (n, h, t) in enumerate(br):
    x = 0.8 + k * 3.0
    y = 5.0 - k * 0.75
    box(s, x, y, 2.8, 6.25 - y, BLUE if k % 2 == 0 else TEAL, [(n + "  " + h, 17, True, "FFFFFF"), (t, 12, False, "FFFFFF")],
        anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margin=0.15)

s = new_slide([("h", "Warum Emotionale Intelligenz im Beruf zählt")],
              "Goleman (1998), HBR „What Makes a Leader?“; O'Boyle et al. (2011), J. Organiz. Behavior; Seligman & Schulman (1986)",
              body_h=0.6)
stats = [("2×", "so wichtig wie IQ und Fachwissen zusammen – so Golemans Analyse von Kompetenzmodellen für Führungsrollen"),
         ("ρ ≈ .30", "Zusammenhang von EI und Arbeitsleistung in einer Metaanalyse – auch über IQ und Persönlichkeit hinaus"),
         ("+37 %", "mehr Abschlüsse erzielten optimistische Versicherungsvertreter in den ersten zwei Jahren")]
for k, (big, t) in enumerate(stats):
    x = 0.8 + k * 3.95
    box(s, x, 2.2, 3.7, 3.4, GREY, [(big, 44, True, BLUE if k != 1 else TEAL), ("", 8, False, DARK), (t, 14, False, DARK)],
        anchor=MSO_ANCHOR.TOP, margin=0.25)
box(s, 0.8, 5.8, 11.7, 0.5, None, [("Aber: Die Effektstärken schwanken je nach Modell und Messmethode – EI ist ein Erfolgsfaktor, keine Garantie.", 14, False, "4A5561")],
    align=PP_ALIGN.LEFT, margin=0.02)

s = new_slide([("h", "Mythen und kritische Einordnung")], "Joseph & Newman (2010); Antonakis et al. (2009); eigene Darstellung", body_h=0.6)
table(s, 0.8, 2.1, 11.7, [
    ["Mythos", "Realität"],
    ["„EI ist angeboren.“", "Temperament ist teils angelegt – Wahrnehmung und Regulation lassen sich aber trainieren."],
    ["„Emotional intelligent = immer nett.“", "EI heißt auch: klar Grenzen setzen, Konflikte ansprechen, unbequeme Wahrheiten sagen."],
    ["„Gefühle haben im Job nichts zu suchen.“", "Emotionen beeinflussen jede Entscheidung – die Frage ist nur, ob bewusst oder unbewusst."],
    ["„Hohe EI schützt vor Fehlern.“", "EI kann auch manipulativ eingesetzt werden – Werte und Integrität bleiben entscheidend."],
    ["„EI ist eindeutig messbar.“", "Selbsteinschätzungen überlappen stark mit Persönlichkeit; Leistungstests sind aufwändig."],
], [4.2, 7.5], size=14, rowh=0.68)

s = new_slide([("h", "Basisemotionen und ihre Botschaft"),
               ("p", "Jede Emotion ist eine **Information** – sie zeigt ein Bedürfnis an und bereitet eine Handlung vor.")],
              "Ekman (1992); eigene Darstellung", body_h=1.2)
table(s, 0.8, 2.55, 11.7, [
    ["Emotion", "Botschaft", "Handlungsimpuls", "Im Arbeitsalltag z. B."],
    ["Freude", "„Das ist gut für mich.“", "Annähern, teilen", "Projekterfolg, Anerkennung"],
    ["Ärger/Wut", "„Eine Grenze wurde verletzt.“", "Verteidigen, durchsetzen", "Unfaire Kritik, Blockaden"],
    ["Angst", "„Gefahr droht.“", "Schützen, vermeiden", "Reorganisation, Präsentation"],
    ["Trauer", "„Ich habe etwas verloren.“", "Innehalten, Unterstützung suchen", "Abschied von Kollegen"],
    ["Ekel", "„Das ist unverträglich.“", "Abwenden, ablehnen", "Unethisches Verhalten"],
    ["Überraschung", "„Das ist neu.“", "Aufmerksamkeit fokussieren", "Unerwartete Zahlen"],
    ["Verachtung", "„Ich stehe darüber.“", "Abwerten, distanzieren", "Warnsignal für Beziehungen"],
], [2.2, 3.4, 3.3, 2.8], size=13, rowh=0.44)

s = new_slide([("h", "Der „Amygdala-Hijack“: Wenn Emotionen übernehmen")],
              "LeDoux (1996); Goleman (1995)", body_h=0.6)
box(s, 0.8, 3.4, 2.1, 1.3, GREY, [("Reiz", 18, True, BLUE), ("z. B. scharfe E-Mail", 12, False, DARK)])
box(s, 3.4, 3.4, 2.1, 1.3, BLUE, [("Thalamus", 18, True, "FFFFFF"), ("Schaltzentrale", 12, False, "E6EEF7")])
arrow(s, 2.95, 3.85, 0.4, 0.4)
arrow(s, 5.6, 2.6, 1.4, 0.4)
arrow(s, 5.6, 5.1, 1.4, 0.4)
box(s, 7.1, 2.1, 5.4, 1.5, "B3261E", [("Schnellweg → Amygdala", 17, True, "FFFFFF"),
                                     ("ca. 12 Millisekunden · unbewusst · Kampf, Flucht, Erstarren", 13, False, "FFFFFF")])
box(s, 7.1, 4.6, 5.4, 1.5, TEAL, [("Langsamweg → Präfrontaler Kortex", 17, True, "FFFFFF"),
                                  ("deutlich langsamer · bewusst · prüfen, abwägen, entscheiden", 13, False, "FFFFFF")])
box(s, 0.8, 5.25, 4.7, 1.0, None, [("Ziel ist nicht, die Amygdala abzuschalten – sondern **Zeit** für den Langsamweg zu gewinnen.", 14, False, DARK)],
    align=PP_ALIGN.LEFT, margin=0.02)

exercise("Mein emotionaler Fingerabdruck", "Welche Emotionen erlebe ich im Job am häufigsten – und was lösen sie aus?",
         ["Denken Sie an die letzten zwei Arbeitswochen.",
          "Notieren Sie die drei häufigsten Emotionen und je eine typische Situation.",
          "Wie haben Sie jeweils reagiert – und mit welcher Wirkung?",
          "Teilen Sie eine Erkenntnis im Chat."],
         "10 Minuten", "Einzelarbeit, danach Chat")

# ---- Modul 2
divider(2, "Das Goleman-Modell", ["Aus welchen Kompetenzen besteht EI?", "Welche betreffen mich – welche andere?",
                                   "Wo liegen meine Stärken und Entwicklungsfelder?"])

s = new_slide([("h", "Die fünf Dimensionen nach Goleman")], "Goleman (1998): Working with Emotional Intelligence", body_h=0.6)
dims = [("Selbst\u00adwahrnehmung", "Eigene Gefühle, Stärken und Grenzen kennen"),
        ("Selbst\u00adregulation", "Impulse steuern, besonnen handeln"),
        ("Motivation", "Aus innerem Antrieb Ziele verfolgen"),
        ("Empathie", "Gefühle und Perspektiven anderer verstehen"),
        ("Soziale Kompetenz", "Beziehungen gestalten, führen, überzeugen")]
box(s, 0.8, 2.1, 7.1, 0.5, LBLUE, [("Persönliche Kompetenz – Umgang mit sich selbst", 14, True, BLUE)], margin=0.02)
box(s, 8.1, 2.1, 4.4, 0.5, LTEAL, [("Soziale Kompetenz – Umgang mit anderen", 14, True, TEAL)], margin=0.02)
for k, (h, t) in enumerate(dims):
    x = 0.8 + k * 2.35
    col = BLUE if k < 3 else TEAL
    box(s, x + 0.35, 2.9, 1.5, 1.5, col, [(str(k + 1), 30, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, x, 4.6, 2.2, 1.7, GREY, [(h.replace("\n", ""), 15, True, col), (t, 12, False, DARK)],
        anchor=MSO_ANCHOR.TOP, margin=0.1)

s = new_slide([("h", "Das Kompetenzmodell nach Goleman & Boyatzis (ESCI)")],
              "Goleman & Boyatzis (2017), HBR; Emotional and Social Competency Inventory (ESCI)", body_h=0.6)
quad = [("Selbstwahrnehmung", BLUE, ["Emotionale Selbstwahrnehmung"]),
        ("Soziales Bewusstsein", TEAL, ["Empathie", "Organisationsbewusstsein"]),
        ("Selbstmanagement", BLUE, ["Emotionale Selbstkontrolle", "Anpassungsfähigkeit", "Leistungsorientierung", "Positive Grundhaltung"]),
        ("Beziehungsmanagement", TEAL, ["Einfluss", "Coaching und Mentoring", "Konfliktmanagement", "Teamarbeit", "Inspirierende Führung"])]
box(s, 2.3, 2.05, 5.0, 0.4, None, [("SELBST", 13, True, "4A5561")], margin=0)
box(s, 7.5, 2.05, 5.0, 0.4, None, [("ANDERE", 13, True, "4A5561")], margin=0)
box(s, 0.8, 2.5, 1.4, 1.75, None, [("Wahr-\nnehmen", 13, True, "4A5561")], margin=0)
box(s, 0.8, 4.35, 1.4, 1.95, None, [("Steuern", 13, True, "4A5561")], margin=0)
for k, (h, col, items) in enumerate(quad):
    x = 2.3 + (k % 2) * 5.2
    y = 2.5 if k < 2 else 4.35
    hgt = 1.75 if k < 2 else 1.95
    box(s, x, y, 5.0, hgt, LBLUE if col == BLUE else LTEAL,
        [(h, 16, True, col), (" · ".join(items), 13, False, DARK)], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.18)

# ---- Modul 3 Selbstwahrnehmung
divider(3, "Selbstwahrnehmung", ["Was fühle ich gerade – und woran merke ich das?", "Was sind meine emotionalen Auslöser?",
                                  "Wie sehen mich andere?"])

s = new_slide([("h", "Selbstwahrnehmung – die Basis aller EI-Kompetenzen"),
               ("p", "Selbstwahrnehmung ist die Fähigkeit, die **eigenen Emotionen, Stärken, Schwächen, Werte und Wirkungen** "
                     "zu erkennen – im Moment und im Rückblick."),
               ("sh", "Woran man hohe Selbstwahrnehmung erkennt"),
               ("b", "Realistische Selbsteinschätzung:: kennt Stärken und Grenzen, spricht offen darüber"),
               ("b", "Selbstbewusstsein:: sicheres Auftreten ohne Arroganz"),
               ("b", "Offenheit für Feedback:: sucht aktiv nach Rückmeldung"),
               ("b", "Humor über sich selbst:: kann eigene Fehler mit Gelassenheit betrachten"),
               ("sh", "Warum das wichtig ist"),
               ("p", "Nach Studien von Tasha Eurich halten sich rund **95 %** der Menschen für selbstreflektiert – "
                     "tatsächlich trifft das nur auf etwa **10–15 %** zu.")],
              "Goleman (1998); Eurich (2018): Insight / HBR „What Self-Awareness Really Is“")

s = new_slide([("h", "Emotionale Granularität: Gefühle genau benennen"),
               ("p", "Je **präziser** wir Gefühle benennen können, desto besser können wir sie regulieren (Lisa Feldman Barrett). "
                     "Aus „mir geht's schlecht“ wird z. B. „ich bin enttäuscht“ oder „ich bin überfordert“.")],
              "Barrett (2017): How Emotions Are Made; Kashdan, Barrett & McKnight (2015)", body_h=1.6)
fam = [("Ärger", ["gereizt", "frustriert", "verärgert", "empört", "wütend"]),
       ("Angst", ["unsicher", "besorgt", "nervös", "beunruhigt", "panisch"]),
       ("Trauer", ["enttäuscht", "niedergeschlagen", "einsam", "entmutigt", "verletzt"]),
       ("Freude", ["zufrieden", "erleichtert", "dankbar", "stolz", "begeistert"])]
for k, (h, ws) in enumerate(fam):
    x = 0.8 + k * 2.95
    box(s, x, 3.2, 2.75, 0.6, BLUE if k % 2 == 0 else TEAL, [(h, 17, True, "FFFFFF")])
    for j, w in enumerate(ws):
        box(s, x, 3.9 + j * 0.47, 2.75, 0.4, GREY, [(w, 13, False, DARK)], margin=0.02)

s = new_slide([("h", "Körpersignale als Frühwarnsystem"),
               ("p", "Emotionen zeigen sich im Körper, **bevor** sie uns bewusst werden. Wer diese Signale kennt, gewinnt Zeit zum Reagieren.")],
              "Damasio (1994): Descartes' Irrtum; Nummenmaa et al. (2014), PNAS", body_h=1.3)
sig = [("Kopf", "Druck, Hitze, Gedankenrasen"), ("Kiefer & Nacken", "Anspannung, Zähne zusammenbeißen"),
       ("Brust", "Enge, flache Atmung, Herzklopfen"), ("Bauch", "Flaues Gefühl, „Kloß im Magen“"),
       ("Hände", "Schwitzen, Fäuste ballen, Zittern")]
for k, (h, t) in enumerate(sig):
    y = 2.85 + k * 0.7
    box(s, 0.8, y, 2.6, 0.58, TEAL, [(h, 14, True, "FFFFFF")])
    box(s, 3.55, y, 4.3, 0.58, GREY, [(t, 13, False, DARK)], align=PP_ALIGN.LEFT)
box(s, 8.3, 2.85, 4.2, 3.4, LBLUE, [("Mini-Body-Scan (60 Sek.)", 16, True, BLUE), ("", 6, False, DARK),
                                    ("1. Kurz innehalten, Augen schließen", 13, False, DARK),
                                    ("2. Aufmerksamkeit vom Kopf bis zu den Füßen wandern lassen", 13, False, DARK),
                                    ("3. Wo spüre ich Spannung, Wärme, Enge?", 13, False, DARK),
                                    ("4. Welches Gefühl könnte dahinterstehen?", 13, False, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Das Johari-Fenster: Selbstbild und Fremdbild")], "Luft & Ingham (1955)", body_h=0.6)
jo = [("Öffentliche Person", "Mir bekannt – anderen bekannt", BLUE),
      ("Blinder Fleck", "Mir unbekannt – anderen bekannt", TEAL),
      ("Privatperson", "Mir bekannt – anderen unbekannt", TEAL),
      ("Unbekanntes", "Mir unbekannt – anderen unbekannt", "6B7B8C")]
for k, (h, t, c) in enumerate(jo):
    x = 1.3 + (k % 2) * 3.55
    y = 2.35 + (k // 2) * 1.95
    box(s, x, y, 3.45, 1.85, c, [(h, 17, True, "FFFFFF"), (t, 12, False, "FFFFFF")], shape=MSO_SHAPE.RECTANGLE)
box(s, 1.3, 1.95, 7.0, 0.35, None, [("bekannt        ←  Selbst  →        unbekannt", 12, True, "4A5561")], margin=0)
box(s, 8.95, 2.35, 3.55, 3.75, LBLUE, [("So wächst die „Arena“", 16, True, BLUE), ("", 6, False, DARK),
                                      ("Feedback einholen → blinder Fleck wird kleiner", 13, False, DARK), ("", 6, False, DARK),
                                      ("Sich öffnen → Privatbereich wird kleiner", 13, False, DARK), ("", 6, False, DARK),
                                      ("Größere Arena = mehr Vertrauen und weniger Missverständnisse", 13, True, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Auslöser erkennen: Das ABC-Modell"),
               ("p", "Nicht die Situation selbst, sondern unsere **Bewertung** bestimmt, was wir fühlen (Albert Ellis).")],
              "Ellis (1962): Reason and Emotion in Psychotherapy", body_h=1.2)
abc = [("A", "Auslöser", "Kollege unterbricht mich im Meeting"),
       ("B", "Bewertung", "„Er nimmt mich nicht ernst.“"),
       ("C", "Consequence", "Ärger, Rückzug, spitze Bemerkung")]
for k, (l, h, t) in enumerate(abc):
    x = 0.8 + k * 4.05
    box(s, x, 2.75, 3.5, 2.2, BLUE if k != 1 else TEAL, [(l, 36, True, "FFFFFF"), (h, 16, True, "FFFFFF"), (t, 13, False, "FFFFFF")],
        anchor=MSO_ANCHOR.TOP, margin=0.15)
    if k < 2:
        arrow(s, x + 3.58, 3.65, 0.4, 0.4)
box(s, 0.8, 5.2, 11.7, 1.05, LTEAL, [("Alternative Bewertung: „Er ist unter Zeitdruck und will schnell zum Punkt.“ → Gelassenheit, "
                                      "sachliche Bitte: „Lass mich kurz ausreden, dann bist du dran.“", 14, False, DARK)],
    align=PP_ALIGN.LEFT, margin=0.2)

exercise("Feedback zum blinden Fleck", "Das eigene Selbstbild mit dem Fremdbild abgleichen.",
         ["Wählen Sie drei Eigenschaften, die Sie im Umgang mit Emotionen beschreiben.",
          "Tauschen Sie sich im Breakout-Raum (3 Personen) über eine typische Arbeitssituation aus.",
          "Die anderen geben Rückmeldung: „Auf mich wirkst du …“",
          "Notieren Sie: Was bestätigt sich – was überrascht mich?"],
         "20 Minuten", "Breakout-Räume à 3 Personen",
         notes="Feedbackregeln vorher kurz wiederholen: beschreiben statt bewerten, Ich-Botschaften, konkret.")

# ---- Modul 4 Selbstregulation
divider(4, "Selbstregulation", ["Wie behalte ich unter Druck einen klaren Kopf?", "Welche Strategien helfen – welche schaden?",
                                 "Was kann ich in akuten Situationen sofort tun?"])

s = new_slide([("h", "Selbstregulation heißt nicht Unterdrücken"),
               ("p", "Selbstregulation ist die Fähigkeit, **Impulse und Stimmungen zu steuern** und erst zu denken, dann zu handeln."),
               ("sh", "Merkmale"),
               ("b", "Selbstkontrolle:: bleibt auch in Konflikten sachlich"),
               ("b", "Anpassungsfähigkeit:: geht flexibel mit Veränderung und Unsicherheit um"),
               ("b", "Vertrauenswürdigkeit:: handelt berechenbar und im Einklang mit eigenen Werten"),
               ("b", "Verantwortung:: steht zu Fehlern, statt Schuld abzuschieben"),
               ("sh", "Wichtig"),
               ("p", "Gefühle zu unterdrücken kostet Energie, belastet Beziehungen und macht Gefühle oft stärker. "
                     "Ziel ist ein **bewusster Umgang** – nicht emotionslose Kontrolle.")],
              "Goleman (1998); Gross & John (2003)")

s = new_slide([("h", "Das Prozessmodell der Emotionsregulation"),
               ("p", "James Gross unterscheidet fünf Ansatzpunkte – je früher, desto weniger Kraft kostet die Regulation.")],
              "Gross (1998, 2015): The Emerging Field of Emotion Regulation", body_h=1.2)
gp = [("Situation\nauswählen", "Kritische Situationen bewusst meiden oder aufsuchen"),
      ("Situation\nverändern", "Rahmen gestalten: Agenda, Ort, Zeitpunkt"),
      ("Fokus\nlenken", "Aufmerksamkeit auf Lösungen statt auf Ärgernisse richten"),
      ("Neu\nbewerten", "Situation anders interpretieren (Reframing)"),
      ("Reaktion\nsteuern", "Atmen, Pause, Ausdruck bewusst dosieren")]
for k, (h, t) in enumerate(gp):
    x = 0.8 + k * 2.37
    chev(s, x, 2.8, 2.3, 1.2, BLUE if k < 4 else TEAL, [(h.replace("\n", " "), 13, True, "FFFFFF")], first=(k == 0))
    box(s, x, 4.2, 2.1, 2.0, GREY, [(t, 13, False, DARK)], anchor=MSO_ANCHOR.TOP, margin=0.12)
box(s, 0.8, 6.3, 9.2, 0.35, None, [("früh (vor der Emotion)  ─────────────────────────►  spät (während der Emotion)", 12, True, "4A5561")],
    align=PP_ALIGN.LEFT, margin=0)

s = new_slide([("h", "Neubewerten schlägt Unterdrücken")], "Gross & John (2003), J. Personality & Social Psychology", body_h=0.6)
box(s, 0.8, 2.15, 5.7, 0.8, TEAL, [("Kognitive Neubewertung", 20, True, "FFFFFF")])
box(s, 6.8, 2.15, 5.7, 0.8, "8A95A1", [("Unterdrückung", 20, True, "FFFFFF")])
box(s, 0.8, 3.05, 5.7, 3.2, LTEAL, [("„Was könnte es noch bedeuten?“", 15, True, TEAL), ("", 6, False, DARK),
                                   ("+ mehr positive Emotionen", 14, False, DARK), ("+ bessere Beziehungen", 14, False, DARK),
                                   ("+ höheres Wohlbefinden", 14, False, DARK), ("+ Gedächtnis und Denken bleiben frei", 14, False, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.25)
box(s, 6.8, 3.05, 5.7, 3.2, GREY, [("„Nur nichts anmerken lassen!“", 15, True, "4A5561"), ("", 6, False, DARK),
                                  ("– innere Erregung bleibt bestehen", 14, False, DARK), ("– wirkt auf andere unecht, distanziert", 14, False, DARK),
                                  ("– geringeres Wohlbefinden", 14, False, DARK), ("– kostet kognitive Ressourcen", 14, False, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.25)

s = new_slide([("h", "„Name it to tame it“ – Gefühle benennen"),
               ("p", "Schon das **Benennen** eines Gefühls („Ich bin gerade verärgert“) dämpft nachweislich die Aktivität der Amygdala "
                     "und stärkt den präfrontalen Kortex (Affect Labeling)."),
               ("sh", "So geht's in der Praxis"),
               ("n", "Innehalten und wahrnehmen: „Was passiert gerade in mir?“"),
               ("n", "Gefühl möglichst präzise benennen – still oder laut"),
               ("n", "Intensität einschätzen: 1 (leicht) bis 10 (sehr stark)"),
               ("n", "Bedürfnis dahinter erkennen: „Was brauche ich jetzt?“"),
               ("sh", "Beispiel im Meeting"),
               ("i", "„Ich merke, dass mich das gerade ärgert. Ich brauche einen Moment, bevor ich antworte.“")],
              "Lieberman et al. (2007), Psychological Science; Siegel (2010): Mindsight")

s = new_slide([("h", "Soforthilfe in emotional aufgeladenen Situationen")],
              "Bolte Taylor (2008); Zaccaro et al. (2018), Frontiers in Human Neuroscience; eigene Darstellung", body_h=0.6)
tools = [("STOP", "**S**topp – **T**ief durchatmen – **O**bservieren, was los ist – **P**ositiv und bewusst weitermachen"),
         ("Box-Atmung", "4 Sek. einatmen – 4 halten – 4 ausatmen – 4 halten. 3–4 Runden beruhigen das Nervensystem."),
         ("90-Sekunden-Regel", "Die chemische Welle einer Emotion ebbt nach ca. 90 Sekunden ab – wenn wir sie nicht mit Gedanken nähren."),
         ("Perspektivwechsel", "„Wie wichtig ist das in 10 Tagen, 10 Monaten, 10 Jahren?“"),
         ("Bewusste Pause", "„Ich melde mich dazu heute Nachmittag.“ – Online: Kamera kurz aus, Wasser holen."),
         ("Bewegung", "Aufstehen, kurzer Gang, Schultern lockern – Stresshormone werden abgebaut.")]
for k, (h, t) in enumerate(tools):
    x = 0.8 + (k % 3) * 3.95
    y = 2.15 + (k // 3) * 2.1
    box(s, x, y, 3.7, 1.95, GREY, [(h, 17, True, BLUE if k % 2 == 0 else TEAL), (t, 13, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.18)

s = new_slide([("h", "Stress, Emotionen und Resilienz")], "Selye (1974); Lazarus & Folkman (1984); Schneider et al. (2013)", body_h=0.6)
box(s, 0.8, 2.15, 5.6, 4.1, LBLUE, [("Stress entsteht im Kopf", 18, True, BLUE), ("", 6, False, DARK),
                                    ("Nach Lazarus entscheidet unsere **Bewertung**, ob eine Anforderung als Herausforderung (Eustress) "
                                     "oder als Bedrohung (Distress) erlebt wird:", 14, False, DARK), ("", 6, False, DARK),
                                    ("1. Ist die Situation bedrohlich?", 14, False, DARK),
                                    ("2. Habe ich genug Ressourcen, sie zu bewältigen?", 14, False, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.25)
box(s, 6.8, 2.15, 5.7, 4.1, LTEAL, [("EI stärkt Resilienz", 18, True, TEAL), ("", 6, False, DARK),
                                    ("Menschen mit hoher EI bewerten Stressoren eher als Herausforderung, zeigen weniger "
                                     "physiologische Stressreaktion und erholen sich schneller.", 14, False, DARK), ("", 6, False, DARK),
                                    ("Hebel: Selbstwahrnehmung + Neubewertung + soziale Unterstützung", 14, True, DARK)],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.25)

exercise("Reframing", "Belastende Situationen bewusst neu bewerten.",
         ["Jede Person schildert kurz eine aktuelle Situation, die Ärger oder Druck auslöst.",
          "Die Gruppe sammelt mindestens drei alternative Deutungen.",
          "Die erzählende Person wählt die hilfreichste Deutung aus.",
          "Wie verändert sich das Gefühl (Skala 1–10) vorher/nachher?"],
         "25 Minuten", "Breakout-Räume à 3–4 Personen",
         notes="Leitfragen für Reframing: Was könnte die andere Person gute Absichten haben? Welche Chance steckt darin? Was würde ein Mentor sagen?")

# ---- Modul 5 Motivation
divider(5, "Motivation", ["Was treibt mich wirklich an?", "Wie bleibe ich bei Rückschlägen dran?",
                           "Wie beeinflusst meine Denkweise meinen Erfolg?"])

s = new_slide([("h", "Intrinsische Motivation: Die Selbstbestimmungstheorie"),
               ("p", "Nachhaltige Motivation entsteht, wenn drei psychologische **Grundbedürfnisse** erfüllt sind.")],
              "Deci & Ryan (1985, 2000): Self-Determination Theory", body_h=1.2)
sdt = [("Autonomie", "Selbst entscheiden, wie ich arbeite", "Handlungsspielräume geben, Ziele statt Wege vorgeben"),
       ("Kompetenz", "Sich wirksam und fähig erleben", "Passende Herausforderungen, Feedback, Lernchancen"),
       ("Zugehörigkeit", "Sich verbunden und wertgeschätzt fühlen", "Vertrauen, Teamgeist, echtes Interesse")]
for k, (h, t, f) in enumerate(sdt):
    x = 0.8 + k * 3.95
    box(s, x + 0.95, 2.7, 1.8, 1.8, BLUE if k != 1 else TEAL, [(h, 15, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0.02)
    box(s, x, 4.7, 3.7, 1.55, GREY, [(t, 14, True, DARK), ("Führung: " + f, 12, False, "4A5561")], anchor=MSO_ANCHOR.TOP, margin=0.15)

s = new_slide([("h", "Optimismus: Wie wir Rückschläge erklären"),
               ("p", "Martin Seligman zeigt: Entscheidend ist der **Erklärungsstil** – wie wir uns Misserfolge erklären.")],
              "Seligman (1990): Learned Optimism; Seligman & Schulman (1986)", body_h=1.2)
table(s, 0.8, 2.7, 11.7, [
    ["Dimension", "Pessimistischer Stil", "Optimistischer Stil"],
    ["Dauer", "„Das wird immer so sein.“", "„Das ist diesmal passiert.“"],
    ["Reichweite", "„Ich kann einfach nicht verkaufen.“", "„Dieser Kunde hatte gerade kein Budget.“"],
    ["Person", "„Es liegt nur an mir.“", "„Mehrere Faktoren spielten eine Rolle – was kann ich beeinflussen?“"],
], [2.5, 4.4, 4.8], size=14, rowh=0.62)
box(s, 0.8, 5.4, 11.7, 0.85, LTEAL, [("Praxisbeleg: Bei MetLife verkauften optimistische Versicherungsvertreter in zwei Jahren rund 37 % mehr als pessimistische.", 14, True, BLUE)],
    margin=0.15)

s = new_slide([("h", "Belohnungsaufschub und Selbstdisziplin"),
               ("p", "Im **Marshmallow-Test** (Walter Mischel) konnten Kinder zwischen einer Süßigkeit sofort oder zwei Süßigkeiten "
                     "nach Wartezeit wählen. Kinder, die warten konnten, zeigten später im Schnitt bessere Schulleistungen."),
               ("sh", "Kritische Einordnung"),
               ("p", "Eine große Replikation (Watts et al., 2018) fand einen **deutlich kleineren Effekt** – der soziale Hintergrund "
                     "erklärt viel. Selbstkontrolle ist also auch eine Frage von **Vertrauen und Umfeld**."),
               ("sh", "Was hilft Erwachsenen beim Belohnungsaufschub?"),
               ("b", "Wenn-dann-Pläne:: „Wenn die E-Mail mich ärgert, dann antworte ich erst nach dem Mittag.“"),
               ("b", "Versuchungen aus dem Blickfeld:: Benachrichtigungen aus, Fokuszeiten blocken"),
               ("b", "Das „Warum“ sichtbar machen:: Ziel mit persönlichem Sinn verknüpfen")],
              "Mischel et al. (1989), Science; Watts, Duncan & Quan (2018), Psychological Science; Gollwitzer (1999)")

s = new_slide([("h", "Growth Mindset: Fähigkeiten sind entwickelbar")], "Dweck (2006): Mindset", body_h=0.6)
box(s, 0.8, 2.15, 5.7, 0.75, "8A95A1", [("Fixed Mindset", 19, True, "FFFFFF")])
box(s, 6.8, 2.15, 5.7, 0.75, TEAL, [("Growth Mindset", 19, True, "FFFFFF")])
pairs = [("„Das kann ich nicht.“", "„Das kann ich noch nicht.“"),
         ("Fehler sind peinlich.", "Fehler sind Lernchancen."),
         ("Kritik ist ein Angriff.", "Kritik ist wertvolle Information."),
         ("Der Erfolg anderer bedroht mich.", "Der Erfolg anderer inspiriert mich."),
         ("Anstrengung zeigt fehlendes Talent.", "Anstrengung ist der Weg zur Meisterschaft.")]
for k, (a, b) in enumerate(pairs):
    y = 3.05 + k * 0.63
    box(s, 0.8, y, 5.7, 0.53, GREY, [(a, 14, False, DARK)], margin=0.03)
    box(s, 6.8, y, 5.7, 0.53, LTEAL, [(b, 14, True, DARK)], margin=0.03)

exercise("Mein „Warum“", "Die eigene Motivation klären und als Energiequelle nutzen.",
         ["Was gibt mir in meiner Arbeit am meisten Energie?",
          "Wann war ich zuletzt richtig stolz auf meine Arbeit – und warum?",
          "Welche meiner Werte zeigen sich darin?",
          "Formulieren Sie einen Satz: „Ich mache meine Arbeit, weil …“"],
         "10 Minuten", "Einzelarbeit, freiwillig im Plenum teilen")

# ---- Modul 6 Empathie
divider(6, "Empathie", ["Was ist Empathie – und was nicht?", "Wie höre ich wirklich zu?",
                         "Wie erkenne ich Emotionen – auch in Videokonferenzen?"])

s = new_slide([("h", "Drei Formen der Empathie")], "Goleman (2013): Focus; Ekman (2003)", body_h=0.6)
emp = [("Kognitive Empathie", "„Ich verstehe, wie du denkst.“", "Perspektive des anderen nachvollziehen – wichtig für Verhandlung und Führung. Allein genutzt: Gefahr der Manipulation."),
       ("Emotionale Empathie", "„Ich fühle, was du fühlst.“", "Gefühle des anderen mitschwingen lassen – schafft Nähe. Übermaß: Gefahr der emotionalen Erschöpfung."),
       ("Empathische Anteilnahme", "„Ich will dir helfen.“", "Verstehen und Mitfühlen führen zu hilfreichem Handeln – die Königsdisziplin.")]
for k, (h, q, t) in enumerate(emp):
    x = 0.8 + k * 3.95
    box(s, x, 2.15, 3.7, 1.2, [BLUE, TEAL, "2E6B5E"][k], [(h, 17, True, "FFFFFF"), (q, 13, False, "FFFFFF")], margin=0.1)
    box(s, x, 3.45, 3.7, 2.8, GREY, [(t, 14, False, DARK)], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Aktives Zuhören"),
               ("p", "Wir hören meist mit der Absicht zu antworten – nicht mit der Absicht zu **verstehen**.")],
              "Rogers (1951); eigene Darstellung", body_h=1.2)
az = [("Aufmerksam sein", "Blickkontakt, zugewandte Haltung, nicht unterbrechen, Handy weg", "„Hm“, Nicken"),
      ("Paraphrasieren", "Inhalt mit eigenen Worten wiedergeben", "„Sie meinen also, dass …“"),
      ("Verbalisieren", "Das wahrgenommene Gefühl ansprechen", "„Das klingt, als wären Sie enttäuscht.“"),
      ("Nachfragen", "Offene Fragen zur Klärung stellen", "„Was genau ist Ihnen dabei wichtig?“")]
for k, (h, t, e) in enumerate(az):
    y = 2.7 + k * 0.9
    box(s, 0.8, y, 0.7, 0.7, BLUE if k % 2 == 0 else TEAL, [(str(k + 1), 17, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, 1.65, y, 2.8, 0.7, None, [(h, 16, True, BLUE)], align=PP_ALIGN.LEFT, margin=0.02)
    box(s, 4.5, y, 4.3, 0.7, None, [(t, 13, False, DARK)], align=PP_ALIGN.LEFT, margin=0.02)
    box(s, 8.95, y, 3.55, 0.7, LBLUE, [(e, 13, False, DARK)], align=PP_ALIGN.LEFT, margin=0.1)

s = new_slide([("h", "Das Vier-Seiten-Modell nach Schulz von Thun"),
               ("p", "Jede Nachricht hat vier Seiten – und wir hören sie mit „vier Ohren“. Emotional intelligente Menschen "
                     "hören bewusst auf das **Beziehungs- und Appell-Ohr**.")],
              "Schulz von Thun (1981): Miteinander reden 1", body_h=1.5)
box(s, 4.9, 3.35, 3.5, 1.9, GREY, [("„Die Präsentation ist ja immer noch nicht fertig.“", 15, True, DARK)], shape=MSO_SHAPE.OVAL, margin=0.3)
vs = [("Sachinhalt", "Die Präsentation ist nicht fertig.", 0.8, 2.95),
      ("Selbstoffenbarung", "Ich bin unter Druck.", 8.7, 2.95),
      ("Beziehung", "Du bist unzuverlässig.", 0.8, 4.75),
      ("Appell", "Beeil dich!", 8.7, 4.75)]
for k, (h, t, x, y) in enumerate(vs):
    box(s, x, y, 3.8, 1.35, BLUE if k in (0, 3) else TEAL, [(h, 16, True, "FFFFFF"), (t, 13, False, "FFFFFF")], margin=0.1)

s = new_slide([("h", "Emotionen lesen – auch in Videokonferenzen"),
               ("sh", "Worauf achten?"),
               ("b", "Mimik:: Mikroexpressionen, Stirnrunzeln, angespannte Lippen"),
               ("b", "Stimme:: Tempo, Lautstärke, Pausen, Tonfall"),
               ("b", "Sprache:: absolute Wörter („immer“, „nie“), plötzliche Einsilbigkeit"),
               ("b", "Verhalten:: Kamera aus, Schweigen, verspätete Antworten"),
               ("sh", "Besonderheiten online"),
               ("b", "Nonverbale Signale sind stark reduziert – **nachfragen** statt interpretieren"),
               ("b", "Check-ins zu Beginn: „Wie geht es euch gerade – in einem Wort?“"),
               ("b", "Emotionale Zustände explizit machen: „Ich sehe ein paar nachdenkliche Gesichter …“"),
               ("b", "Videokonferenzen ermüden (hohe nonverbale Last) – Pausen einplanen")],
              "Ekman (2003); Bailenson (2021), Technology, Mind, and Behavior")

s = new_slide([("h", "Empathie in virtuellen und hybriden Teams")], "Edmondson (2019); eigene Darstellung", body_h=0.6)
vt = [("Beziehung vor Aufgabe", "5 Minuten informeller Austausch zu Beginn jedes Meetings"),
      ("Regelmäßige 1:1-Gespräche", "Nicht nur Status abfragen – auch Befinden und Belastung"),
      ("Kanal bewusst wählen", "Heikle Themen per Video oder Telefon, nie per Chat oder E-Mail"),
      ("Emojis und Ton", "Schriftliche Nachrichten wirken oft kühler als gemeint – freundlich formulieren"),
      ("Gleiche Bühne für alle", "In hybriden Meetings Remote-Teilnehmende aktiv einbeziehen"),
      ("Wertschätzung sichtbar machen", "Erfolge öffentlich würdigen, z. B. im Team-Kanal")]
for k, (h, t) in enumerate(vt):
    x = 0.8 + (k % 2) * 5.95
    y = 2.15 + (k // 2) * 1.4
    box(s, x, y, 0.55, 0.55, TEAL, [("✓", 16, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, x + 0.7, y - 0.1, 5.0, 1.2, None, [(h, 16, True, BLUE), (t, 13, False, DARK)], align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.02)

exercise("Perspektivwechsel", "Empathie in einer schwierigen Gesprächssituation trainieren.",
         ["Rollen verteilen: Person A (Anliegen), Person B (Zuhörer/in), Person C (Beobachter/in).",
          "A schildert ein echtes Ärgernis aus dem Arbeitsalltag (3 Min.).",
          "B hört aktiv zu: paraphrasieren, Gefühle verbalisieren, nachfragen.",
          "C gibt Feedback: Was hat Verständnis erzeugt? Danach Rollen tauschen."],
         "30 Minuten", "Breakout-Räume à 3 Personen",
         notes="Beobachtungsbogen im Chat bereitstellen: Blickkontakt, Paraphrasen, Gefühlsbenennung, Fragen.")

# ---- Modul 7 Soziale Kompetenz
divider(7, "Soziale Kompetenz", ["Wie gestalte ich tragfähige Beziehungen?", "Wie spreche ich Kritik wertschätzend an?",
                                  "Wie löse ich Konflikte emotional intelligent?"])

s = new_slide([("h", "Beziehungsmanagement: EI in Aktion"),
               ("p", "Soziale Kompetenz bündelt alle anderen EI-Dimensionen: Wer sich selbst kennt und steuert und andere versteht, "
                     "kann **Beziehungen gezielt gestalten**.")],
              "Goleman (1998); Hatfield, Cacioppo & Rapson (1994): Emotional Contagion", body_h=1.4)
bm = [("Überzeugen", "Argumente auf Bedürfnisse des Gegenübers abstimmen"),
      ("Kommunizieren", "Klar, offen und wertschätzend sprechen"),
      ("Konflikte lösen", "Spannungen früh ansprechen, Win-win suchen"),
      ("Zusammenarbeiten", "Vertrauen aufbauen, Wissen teilen"),
      ("Inspirieren", "Begeisterung wecken, Sinn vermitteln"),
      ("Ansteckung nutzen", "Stimmungen übertragen sich – besonders von Führungskräften")]
for k, (h, t) in enumerate(bm):
    x = 0.8 + (k % 3) * 3.95
    y = 3.0 + (k // 3) * 1.65
    box(s, x, y, 3.7, 1.45, LBLUE if (k + k // 3) % 2 == 0 else LTEAL, [(h, 16, True, BLUE), (t, 13, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.18)

s = new_slide([("h", "Gewaltfreie Kommunikation nach Rosenberg"),
               ("p", "Vier Schritte, um Anliegen klar und ohne Vorwurf zu äußern.")],
              "Rosenberg (2001): Gewaltfreie Kommunikation", body_h=1.0)
gfk = [("Beobachtung", "Was genau ist passiert – ohne Bewertung?", "„In den letzten drei Meetings kam der Bericht nach 10 Uhr.“"),
       ("Gefühl", "Was fühle ich dabei?", "„Ich bin beunruhigt …“"),
       ("Bedürfnis", "Welches Bedürfnis steckt dahinter?", "„… weil mir Planungssicherheit wichtig ist.“"),
       ("Bitte", "Konkrete, erfüllbare Bitte", "„Kannst du ihn künftig bis 9 Uhr schicken?“")]
for k, (h, q, e) in enumerate(gfk):
    x = 0.8 + k * 2.97
    chev(s, x, 2.55, 2.85, 1.0, BLUE if k % 2 == 0 else TEAL, [(h, 16, True, "FFFFFF")], first=(k == 0))
    box(s, x, 3.7, 2.7, 1.0, None, [(q, 13, True, DARK)], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05)
    box(s, x, 4.7, 2.7, 1.55, LBLUE, [(e, 13, False, DARK)], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.12)

s = new_slide([("h", "Feedback emotional intelligent geben: Das SBI-Modell"),
               ("p", "Konkretes, beschreibendes Feedback reduziert Abwehr – weil es die **Person nicht angreift**.")],
              "Center for Creative Leadership: Situation-Behavior-Impact", body_h=1.2)
sbi = [("S", "Situation", "Wann und wo?", "„Im Kundentermin gestern …“"),
       ("B", "Behavior (Verhalten)", "Was genau war beobachtbar?", "„… hast du den Kunden zweimal unterbrochen.“"),
       ("I", "Impact (Wirkung)", "Welche Wirkung hatte das?", "„Er wirkte danach verschlossen, und ich hatte Sorge um den Auftrag.“")]
for k, (l, h, q, e) in enumerate(sbi):
    y = 2.7 + k * 1.15
    box(s, 0.8, y, 1.0, 1.0, BLUE if k != 1 else TEAL, [(l, 26, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, 2.0, y, 3.9, 1.0, None, [(h, 16, True, BLUE), (q, 13, False, DARK)], align=PP_ALIGN.LEFT, margin=0.02)
    box(s, 6.1, y + 0.05, 6.4, 0.9, LBLUE, [(e, 14, False, DARK)], align=PP_ALIGN.LEFT, margin=0.15)
box(s, 0.8, 6.15, 11.7, 0.4, None, [("Danach: Pause lassen, Sicht des anderen erfragen, gemeinsam Lösung vereinbaren.", 14, True, TEAL)],
    align=PP_ALIGN.LEFT, margin=0)

s = new_slide([("h", "Konflikte emotional intelligent lösen"),
               ("p", "Das Harvard-Konzept verbindet Sachlichkeit mit Respekt vor den Emotionen der Beteiligten.")],
              "Fisher, Ury & Patton (1981): Das Harvard-Konzept", body_h=1.2)
hv = [("Mensch und Problem trennen", "Hart in der Sache, weich zum Menschen. Emotionen anerkennen, bevor man argumentiert."),
      ("Interessen statt Positionen", "Hinter „Ich will X“ nach dem „Warum?“ fragen – Interessen sind verhandelbar."),
      ("Optionen entwickeln", "Erst Ideen sammeln, dann bewerten. Mehrere Lösungen erhöhen die Chance auf Einigung."),
      ("Objektive Kriterien", "Faire, neutrale Maßstäbe vereinbaren – das nimmt Emotionen aus Machtfragen.")]
for k, (h, t) in enumerate(hv):
    x = 0.8 + (k % 2) * 5.95
    y = 2.65 + (k // 2) * 1.85
    box(s, x, y, 5.75, 1.65, GREY, [(str(k + 1) + ". " + h, 16, True, BLUE if k % 2 == 0 else TEAL), (t, 13, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Psychologische Sicherheit im Team")],
              "Edmondson (1999), Administrative Science Quarterly; Google re:Work – Project Aristotle (2015)", body_h=0.6)
box(s, 0.8, 2.15, 5.3, 4.1, BLUE, [("Psychologische Sicherheit", 20, True, "FFFFFF"), ("", 8, False, "FFFFFF"),
                                  ("Die geteilte Überzeugung, dass man im Team Risiken eingehen kann – Fragen stellen, Fehler zugeben, "
                                   "Kritik äußern –, ohne bloßgestellt zu werden.", 15, False, "FFFFFF"), ("", 8, False, "FFFFFF"),
                                  ("Googles „Project Aristotle“ identifizierte sie als wichtigsten Faktor erfolgreicher Teams.", 14, True, "FFFFFF")],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.3)
ps = ["Eigene Fehler offen ansprechen", "Fragen stellen statt Antworten vorgeben", "Neugierig auf Einwände reagieren",
      "Beiträge wertschätzen – auch wenn sie nicht passen", "Fehler als Lernchance behandeln"]
box(s, 6.5, 2.15, 6.0, 0.5, None, [("Was Führungskräfte konkret tun können", 16, True, BLUE)], align=PP_ALIGN.LEFT, margin=0.02)
for k, t in enumerate(ps):
    y = 2.8 + k * 0.7
    box(s, 6.5, y, 0.45, 0.45, TEAL, [("✓", 13, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, 7.1, y - 0.05, 5.4, 0.55, None, [(t, 14, False, DARK)], align=PP_ALIGN.LEFT, margin=0.02)

# ---- Modul 8 Führung & Vertrieb
divider(8, "EI in Führung und Vertrieb", ["Welcher Führungsstil wirkt wann?", "Warum kaufen Kunden emotional?",
                                           "Wie gehe ich mit emotionalen Einwänden um?"])

s = new_slide([("h", "Sechs Führungsstile nach Goleman und ihre Wirkung auf das Klima")],
              "Goleman (2000), HBR „Leadership That Gets Results“", body_h=0.6)
table(s, 0.8, 2.1, 11.7, [
    ["Stil", "Kernsatz", "Wirkung aufs Klima", "Sinnvoll bei …"],
    ["Visionär", "„Kommt mit mir.“", "stark positiv (+.54)", "Neuausrichtung, fehlende Orientierung"],
    ["Coachend", "„Probier es aus.“", "positiv (+.42)", "Entwicklung langfristiger Stärken"],
    ["Gefühlsorientiert", "„Menschen zuerst.“", "positiv (+.46)", "Konflikte heilen, Motivation in Krisen"],
    ["Demokratisch", "„Was meint ihr?“", "positiv (+.43)", "Konsens und Beteiligung gewinnen"],
    ["Fordernd (Schrittmacher)", "„Mach es wie ich, jetzt.“", "negativ (−.25)", "Kurzfristig mit Top-Team"],
    ["Befehlend", "„Tu, was ich sage.“", "negativ (−.26)", "Echte Krise, Notfall"],
], [3.0, 2.8, 2.5, 3.4], size=13, rowh=0.55)
box(s, 0.8, 6.0, 11.7, 0.45, None, [("Erfolgreiche Führungskräfte beherrschen mehrere Stile und wechseln situativ – das erfordert hohe EI.", 14, True, TEAL)],
    align=PP_ALIGN.LEFT, margin=0)

s = new_slide([("h", "Resonante Führung: Die Stimmung macht den Unterschied"),
               ("p", "Führungskräfte sind **emotionale Taktgeber**: Ihre Stimmung überträgt sich auf das Team und beeinflusst Leistung "
                     "und Klima (emotionale Ansteckung)."),
               ("sh", "Resonante Führungskräfte …"),
               ("b", "sind sich ihrer eigenen Wirkung bewusst"),
               ("b", "vermitteln Hoffnung und Zuversicht – auch in schwierigen Phasen"),
               ("b", "zeigen echtes Interesse an Menschen (Mitgefühl)"),
               ("b", "achten auf ihre eigene Erholung, um Dissonanz zu vermeiden"),
               ("sh", "Reflexionsfrage"),
               ("i", "Mit welcher Stimmung betrete ich morgens das (virtuelle) Büro – und was löse ich damit aus?")],
              "Goleman, Boyatzis & McKee (2002): Primal Leadership; Boyatzis & McKee (2005): Resonant Leadership")

s = new_slide([("h", "Emotionale Intelligenz im Vertrieb"),
               ("p", "Menschen entscheiden **emotional** und begründen rational: Nach Damasios Theorie der somatischen Marker sind "
                     "Emotionen für jede Entscheidung unverzichtbar.")],
              "Damasio (1994): Descartes' Irrtum; eigene Darstellung", body_h=1.4)
vt = [("Vor dem Gespräch", ["Eigene Stimmung prüfen", "Kunden-Situation antizipieren", "Ziel und Haltung klären"]),
      ("Im Gespräch", ["Emotionale Signale lesen", "Bedürfnisse erfragen", "Nutzen emotional erlebbar machen"]),
      ("Nach dem Gespräch", ["Absagen nicht persönlich nehmen", "Optimistisch erklären", "Beziehung weiter pflegen"])]
for k, (h, its) in enumerate(vt):
    x = 0.8 + k * 3.95
    chev(s, x, 3.0, 3.7, 0.7, BLUE if k != 1 else TEAL, [(h, 16, True, "FFFFFF")], first=(k == 0))
    box(s, x, 3.85, 3.5, 2.35, GREY, [("• " + i, 14, False, DARK) for i in its], align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.2)

s = new_slide([("h", "Emotionale Einwände souverän behandeln"),
               ("p", "Hinter vielen Einwänden stecken Gefühle wie **Unsicherheit, Angst vor Fehlentscheidung oder Misstrauen**.")],
              "eigene Darstellung", body_h=1.2)
ew = [("Zuhören", "Einwand ausreden lassen, nicht sofort kontern"),
      ("Gefühl anerkennen", "„Ich verstehe, dass das eine große Investition ist.“"),
      ("Hinterfragen", "„Was genau bereitet Ihnen dabei Sorgen?“"),
      ("Lösung anbieten", "Referenz, Testphase, Garantie – passend zum Gefühl"),
      ("Zustimmung prüfen", "„Ist Ihre Sorge damit ausgeräumt?“")]
for k, (h, t) in enumerate(ew):
    x = 0.8 + k * 2.37
    chev(s, x, 2.75, 2.3, 1.05, BLUE if k % 2 == 0 else TEAL, [(h, 13, True, "FFFFFF")], first=(k == 0))
    box(s, x, 3.95, 2.1, 1.8, GREY, [(t, 13, False, DARK)], anchor=MSO_ANCHOR.TOP, margin=0.12)
box(s, 0.8, 5.9, 11.7, 0.4, None, [("Kein Einwand ohne Anerkennung – erst die Beziehung, dann das Argument.", 14, True, TEAL)],
    align=PP_ALIGN.LEFT, margin=0)

exercise("Praxisfälle", "EI-Werkzeuge in typischen Situationen aus Führung und Vertrieb anwenden.",
         ["Fall A: Ein Mitarbeiter reagiert im Teams-Call gereizt auf eine neue Aufgabe.",
          "Fall B: Eine Kundin droht verärgert mit Vertragskündigung.",
          "Analysieren Sie: Welche Emotionen, welche Bedürfnisse, welche eigenen Trigger?",
          "Spielen Sie das Gespräch kurz durch und präsentieren Sie Ihre drei wichtigsten Erkenntnisse."],
         "30 Minuten", "Breakout-Räume, Präsentation im Plenum",
         notes="Gruppen frei wählen lassen: Führungskräfte Fall A, Vertrieb Fall B.")

# ---- Modul 9 Transfer
divider(9, "Transfer in den Alltag", ["Wie mache ich EI zur Gewohnheit?", "Wo stehe ich heute?",
                                       "Was setze ich ab morgen konkret um?"])

s = new_slide([("h", "EI entwickeln: Kleine Routinen, große Wirkung"),
               ("p", "Neue Gewohnheiten brauchen Zeit: Im Schnitt dauerte es **66 Tage**, bis ein neues Verhalten automatisch ablief "
                     "(Spannweite 18–254 Tage).")],
              "Lally et al. (2010), European Journal of Social Psychology; eigene Darstellung", body_h=1.3)
rt = [("Morgens", "2 Minuten: „Wie geht es mir? Was brauche ich heute?“"),
      ("Vor Meetings", "Kurzer Check: Stimmung, Ziel, gewünschte Wirkung"),
      ("In Stressmomenten", "STOP-Technik oder drei tiefe Atemzüge"),
      ("Mittags", "Echte Pause ohne Bildschirm"),
      ("Abends", "Emotionstagebuch: 1 Situation, 1 Gefühl, 1 Erkenntnis"),
      ("Wöchentlich", "Feedback zu einer konkreten Situation einholen")]
for k, (h, t) in enumerate(rt):
    x = 0.8 + (k % 3) * 3.95
    y = 2.9 + (k // 3) * 1.7
    box(s, x, y, 3.7, 1.5, LBLUE if k % 2 == 0 else LTEAL, [(h, 16, True, BLUE), (t, 13, False, DARK)],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.18)

s = new_slide([("h", "Selbstcheck: Wo stehe ich?"),
               ("p", "Bewerten Sie spontan von **1 (trifft nicht zu)** bis **5 (trifft voll zu)**. Hinweis: Reflexionshilfe, kein validierter Test.")],
              "eigene Darstellung in Anlehnung an Goleman (1998)", body_h=1.2)
table(s, 0.8, 2.6, 11.7, [
    ["Dimension", "Aussage", "1 – 5"],
    ["Selbstwahrnehmung", "Ich kann meine Gefühle in den meisten Situationen genau benennen.", ""],
    ["Selbstwahrnehmung", "Ich kenne die Situationen, die mich zuverlässig aus der Ruhe bringen.", ""],
    ["Selbstregulation", "Auch unter Druck reagiere ich überlegt statt impulsiv.", ""],
    ["Selbstregulation", "Nach Rückschlägen finde ich schnell wieder in meine Balance.", ""],
    ["Motivation", "Ich bleibe auch bei Widerständen an meinen Zielen dran.", ""],
    ["Empathie", "Ich merke schnell, wenn es jemandem im Team nicht gut geht.", ""],
    ["Empathie", "Ich höre zu, um zu verstehen – nicht nur, um zu antworten.", ""],
    ["Soziale Kompetenz", "Ich spreche Konflikte früh und wertschätzend an.", ""],
], [2.8, 7.7, 1.2], size=13, rowh=0.41)

s = new_slide([("h", "Mein persönlicher Transferplan"),
               ("p", "Beantworten Sie die Fragen schriftlich – konkret, realistisch und mit Termin.")],
              "eigene Darstellung", body_h=1.2)
tp = [("Meine wichtigste Erkenntnis heute:", LBLUE),
      ("Diese EI-Kompetenz entwickle ich als Erstes:", LTEAL),
      ("Konkret werde ich ab morgen … (Wenn-dann-Plan)", LBLUE),
      ("Woran erkenne ich in 66 Tagen meinen Fortschritt?", LTEAL),
      ("Wer kann mir dabei Feedback geben?", LBLUE)]
for k, (t, c) in enumerate(tp):
    y = 2.65 + k * 0.72
    box(s, 0.8, y, 11.7, 0.62, c, [(t, 15, True, BLUE)], align=PP_ALIGN.LEFT, margin=0.2)

s = new_slide([("h", "Zusammenfassung: Die wichtigsten Erkenntnisse")], "eigene Darstellung", body_h=0.6)
kt = [("1", "Emotionen sind Informationen", "Sie zeigen Bedürfnisse an – wer sie liest, entscheidet besser."),
      ("2", "Selbstwahrnehmung ist die Basis", "Gefühle präzise benennen und Auslöser kennen."),
      ("3", "Steuern statt unterdrücken", "Neubewerten, benennen, bewusst pausieren."),
      ("4", "Empathie verbindet", "Aktiv zuhören, nachfragen, Gefühle anerkennen."),
      ("5", "Beziehungen bewusst gestalten", "Klare, wertschätzende Kommunikation und psychologische Sicherheit."),
      ("6", "EI ist trainierbar", "Kleine tägliche Routinen führen zu nachhaltiger Veränderung.")]
for k, (n, h, t) in enumerate(kt):
    x = 0.8 + (k % 2) * 5.95
    y = 2.15 + (k // 2) * 1.4
    box(s, x, y, 0.8, 0.8, BLUE if k % 2 == 0 else TEAL, [(n, 20, True, "FFFFFF")], shape=MSO_SHAPE.OVAL, margin=0)
    box(s, x + 0.95, y - 0.1, 4.8, 1.25, None, [(h, 16, True, BLUE), (t, 13, False, DARK)], align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, margin=0.02)

s = new_slide([("h", "Literaturempfehlungen"),
               ("b", "Goleman, D. (1996):: EQ. Emotionale Intelligenz. München: Hanser."),
               ("b", "Goleman, D. (1999):: EQ². Der Erfolgsquotient. München: Hanser."),
               ("b", "Goleman, D., Boyatzis, R. & McKee, A. (2003):: Emotionale Führung. Berlin: Ullstein."),
               ("b", "Barrett, L. F. (2023):: Wie Gefühle entstehen. Hamburg: Rowohlt."),
               ("b", "Bradberry, T. & Greaves, J. (2009):: Emotional Intelligence 2.0. San Diego: TalentSmart."),
               ("b", "Rosenberg, M. B. (2016):: Gewaltfreie Kommunikation. Paderborn: Junfermann."),
               ("b", "Schulz von Thun, F. (1981):: Miteinander reden 1. Reinbek: Rowohlt."),
               ("b", "Dweck, C. (2017):: Selbstbild. Wie unser Denken Erfolge oder Niederlagen bewirkt. München: Piper."),
               ("b", "Eurich, T. (2018):: Insight. New York: Crown Business.")],
              "eigene Zusammenstellung")

# Abschlussfolie ans Ende verschieben
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
end_id = ids[len(tpl_slides) - 1]
# Vorlagen-Inhaltsfolien (2..21) entfernen
for sid in ids[1:len(tpl_slides) - 1]:
    rid = sid.get("{%s}id" % NS["r"])
    prs.part.drop_rel(rid)
    sldIdLst.remove(sid)
sldIdLst.remove(end_id)
sldIdLst.append(end_id)
end_slide.notes_slide.notes_text_frame.text = "Offene Fragen klären, Feedbackrunde (Blitzlicht), Verabschiedung."

prs.save(OUT)
print("Folien:", len(prs.slides))
